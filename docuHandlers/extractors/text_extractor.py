"""Generic text extraction from common document types."""
import os
import re


def extract_text(file_path: str) -> dict:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return _from_pdf(file_path)
    if ext in (".docx",):
        return _from_docx(file_path)
    if ext in (".pptx",):
        return _from_pptx(file_path)
    if ext in (".xlsx", ".csv"):
        return _from_spreadsheet(file_path)
    if ext in (".txt", ".md"):
        return _from_plain(file_path)
    return {"text": "", "source": file_path, "error": f"Unsupported file type: {ext}"}


def _from_pdf(path: str) -> dict:
    import pdfplumber

    chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            chunks.append(page.extract_text() or "")
    text = "\n".join(chunks).strip()
    return {"text": text, "source": path, "has_text_layer": bool(text)}


def _from_docx(path: str) -> dict:
    from docx import Document

    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return {"text": text, "source": path}


def _from_pptx(path: str) -> dict:
    from pptx import Presentation

    prs = Presentation(path)
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                chunks.append(shape.text.strip())
    return {"text": "\n".join(chunks), "source": path}


def _from_spreadsheet(path: str) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        cells = [str(c) for c in row if c is not None]
        if cells:
            rows.append(" | ".join(cells))
    return {"text": "\n".join(rows), "source": path}


def _from_plain(path: str) -> dict:
    with open(path) as f:
        return {"text": f.read(), "source": path}


def extract_key_values(text: str) -> list[dict]:
    """Pull label: value pairs from free text for cross-document comparison."""
    pairs = []
    for line in text.splitlines():
        m = re.match(r"^([A-Za-z][A-Za-z\s/%$()]+?)[:=\s]+(.+)$", line.strip())
        if not m:
            continue
        label = m.group(1).strip()
        value = m.group(2).strip()
        if len(label) > 40:
            continue
        pairs.append({"label": label, "value": value})
    return pairs
