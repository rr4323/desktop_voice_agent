"""Correct a document while preserving paragraph/run formatting."""
import copy
import os
import re

from docx import Document

# Built-in low-risk typo dictionary for demo / safe corrections
DEFAULT_CORRECTIONS = {
    r"\bteh\b": "the",
    r"\brecieve\b": "receive",
    r"\boccured\b": "occurred",
    r"\bseperate\b": "separate",
    r"\bdefinately\b": "definitely",
    r"\baccomodate\b": "accommodate",
    r"\bfinacial\b": "financial",
    r"\bquater\b": "quarter",
}


def _parse_replacements(raw: list[str]) -> list[dict]:
    out = []
    for item in raw:
        if "=" not in item:
            continue
        old, new = item.split("=", 1)
        out.append({"pattern": re.escape(old.strip()), "replacement": new.strip(), "confidence": 1.0})
    return out


def detect_corrections(text: str, extra_replacements: list[str] | None = None) -> list[dict]:
    corrections = []
    for pattern, repl in DEFAULT_CORRECTIONS.items():
        if re.search(pattern, text, re.IGNORECASE):
            corrections.append({"pattern": pattern, "replacement": repl, "confidence": 0.9})
    corrections.extend(_parse_replacements(extra_replacements or []))
    return corrections


def correct_document(
    docx_path: str,
    replacements: list[str] | None = None,
    output_path: str | None = None,
) -> dict:
    doc = Document(docx_path)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    rules = detect_corrections(full_text, replacements)

    applied = []
    for para in doc.paragraphs:
        for run in para.runs:
            original = run.text
            updated = original
            for rule in rules:
                updated = re.sub(rule["pattern"], rule["replacement"], updated, flags=re.IGNORECASE)
            if updated != original:
                applied.append({"before": original, "after": updated, "confidence": rule.get("confidence", 1.0)})
                run.text = updated

    out = output_path or docx_path.replace(".docx", "_corrected.docx")
    doc.save(out)

    return {
        "output_path": out,
        "corrections_detected": rules,
        "corrections_applied": applied,
    }
