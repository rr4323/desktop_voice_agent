"""Transfer extracted information from PDF into a Word (.docx) document."""
import os
import re

from docx import Document
from extractors.pdf_extractor import extract_pdf_metrics_map
from extractors.text_extractor import extract_text, extract_key_values


def transfer_pdf_to_docx(
    source_pdf: str, docx_path: str, output_path: str | None = None
) -> dict:
    lookup = extract_pdf_metrics_map(source_pdf)

    # Supplement with general key-values from text extractor
    source_text = extract_text(source_pdf).get("text", "")
    for pair in extract_key_values(source_text):
        slug = re.sub(r"[^a-z0-9]+", "_", pair["label"].lower()).strip("_")
        lookup.setdefault(slug, pair["value"])

    doc = Document(docx_path)
    replacements = []

    sorted_keys = sorted(lookup.keys(), key=lambda k: len(k), reverse=True)

    def process_paragraph(p):
        full_text = p.text
        if not full_text:
            return
        new_text = full_text

        matched_keys = []
        for key in sorted_keys:
            val = str(lookup[key])
            patterns = [
                re.compile(rf"\{{\{{\s*{re.escape(key)}\s*\}}\}}", re.IGNORECASE),
                re.compile(rf"\b{re.escape(key.upper())}\b"),
            ]
            for pat in patterns:
                if pat.search(new_text):
                    new_text = pat.sub(lambda m, v=val: v, new_text)
                    matched_keys.append({"placeholder": key, "value": val})

        if new_text != full_text:
            replacements.extend(matched_keys)
            replaced_in_run = False
            for r in p.runs:
                r_text = r.text
                if not r_text:
                    continue
                r_new = r_text
                for key in sorted_keys:
                    val = str(lookup[key])
                    pat1 = re.compile(rf"\{{\{{\s*{re.escape(key)}\s*\}}\}}", re.IGNORECASE)
                    pat2 = re.compile(rf"\b{re.escape(key.upper())}\b")
                    if pat1.search(r_new) or pat2.search(r_new):
                        r_new = pat1.sub(lambda m, v=val: v, r_new)
                        r_new = pat2.sub(lambda m, v=val: v, r_new)
                        replaced_in_run = True
                r.text = r_new

            if not replaced_in_run and p.runs:
                p.runs[0].text = new_text
                for r in p.runs[1:]:
                    r.text = ""

    for p in doc.paragraphs:
        process_paragraph(p)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_paragraph(p)

    out = output_path or docx_path.replace(".docx", "_updated.docx")
    doc.save(out)

    return {
        "output_path": out,
        "replacements": replacements,
        "values_available": lookup,
        "source": source_pdf,
    }
