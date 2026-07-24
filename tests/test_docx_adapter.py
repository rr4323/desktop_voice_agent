import os
import tempfile
from part4.adapters.docx_adapter import DocxAdapter

def test_docx_read_write():
    tmp = os.path.join(tempfile.gettempdir(), "test_part4_docx.docx")
    # create a small docx
    from docx import Document
    doc = Document()
    doc.add_paragraph("Para one")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Cell00"
    doc.save(tmp)

    adapter = DocxAdapter()
    read = adapter.read_document(tmp)
    assert any(p["text"] == "Para one" for p in read["paragraphs"]) 

    res = adapter.write_paragraph(tmp, 0, "Para one updated")
    assert res.get("success") is True

    res2 = adapter.write_table_cell(tmp, 0, 0, 0, "Cell00-updated")
    assert res2.get("success") is True

if __name__ == "__main__":
    test_docx_read_write()
    print("Docx adapter tests passed")
