import openpyxl
import shutil

class XLSXAdapter:
    """Read/write spreadsheet cells, formulas, and sheets."""

    def read_cell(self, file_path: str, sheet_name: str, cell_ref: str) -> dict:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        try:
            ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
            val = ws[cell_ref].value
            return {"value": val, "cell": cell_ref, "sheet": ws.title}
        except Exception as e:
            return {"error": str(e), "value": None}
        finally:
            wb.close()

    def write_cell(self, file_path: str, sheet_name: str, cell_ref: str, new_value) -> dict:
        backup = file_path + ".bak"
        shutil.copy2(file_path, backup)

        wb = openpyxl.load_workbook(file_path)
        try:
            ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
            previous = ws[cell_ref].value
            ws[cell_ref] = new_value
            wb.save(file_path)
            return {
                "success": True,
                "cell": cell_ref,
                "sheet": ws.title,
                "previous_value": previous,
                "new_value": new_value,
                "backup": backup
            }
        except Exception as e:
            return {"success": False, "error": str(e)}
        finally:
            wb.close()