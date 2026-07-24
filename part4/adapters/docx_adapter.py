from docx import Document
import shutil

class DocxAdapter:
    """Adapter for .docx files using python-docx.

    Methods:
      - read_document(file_path): returns paragraphs and tables with provenance
      - write_paragraph(file_path, paragraph_index, new_text)
      - write_table_cell(file_path, table_index, row_idx, col_idx, new_text)
    """

    def read_document(self, file_path):
        doc = Document(file_path)
        paragraphs = []
        for i, p in enumerate(doc.paragraphs):
            paragraphs.append({"file": file_path, "type": "paragraph", "index": i, "text": p.text})

        tables = []
        for ti, table in enumerate(doc.tables):
            for r, row in enumerate(table.rows):
                for c, cell in enumerate(row.cells):
                    tables.append({"file": file_path, "type": "table_cell", "table_index": ti, "row": r, "col": c, "text": cell.text})

        return {"paragraphs": paragraphs, "table_cells": tables}

    def write_paragraph(self, file_path, paragraph_index, new_text):
        backup = file_path + ".bak"
        shutil.copy2(file_path, backup)
        doc = Document(file_path)
        try:
            prev = doc.paragraphs[paragraph_index].text
        except Exception:
            return {"success": False, "error": "invalid_index"}
        doc.paragraphs[paragraph_index].text = new_text
        doc.save(file_path)
        return {"file": file_path, "paragraph_index": paragraph_index, "previous": prev, "new": new_text, "success": True, "backup": backup}

    def write_table_cell(self, file_path, table_index, row_idx, col_idx, new_text):
        backup = file_path + ".bak"
        shutil.copy2(file_path, backup)
        doc = Document(file_path)
        try:
            cell = doc.tables[table_index].rows[row_idx].cells[col_idx]
        except Exception:
            return {"success": False, "error": "invalid_index"}
        prev = cell.text
        cell.text = new_text
        doc.save(file_path)
        return {"file": file_path, "table_index": table_index, "row": row_idx, "col": col_idx, "previous": prev, "new": new_text, "success": True, "backup": backup}
