from pptx import Presentation
import shutil

class PPTXAdapter:
    """Read/write adapter for presentation slides preserving formatting."""

    def read_slides(self, file_path: str) -> list:
        prs = Presentation(file_path)
        out = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text)
            out.append({"file": file_path, "slide": i, "text": "\n".join(texts)})
        return out

    def write_slide_text(self, file_path: str, slide_index: int, placeholder_name: str, new_text: str) -> dict:
        backup = file_path + ".bak"
        shutil.copy2(file_path, backup)
        prs = Presentation(file_path)

        try:
            slide = prs.slides[slide_index - 1]
        except IndexError:
            return {"success": False, "error": "invalid_slide_index"}

        target_shape = None
        for shape in slide.shapes:
            if shape.name == placeholder_name or (shape.is_placeholder and shape.placeholder_format.type == placeholder_name):
                target_shape = shape
                break

        if not target_shape or not target_shape.has_text_frame:
            return {"success": False, "error": "target_not_found"}

        previous = target_shape.text
        text_frame = target_shape.text_frame
        # Replace text inside runs to preserve paragraph styling
        if text_frame.paragraphs and text_frame.paragraphs[0].runs:
            text_frame.paragraphs[0].runs[0].text = new_text
            for run in text_frame.paragraphs[0].runs[1:]:
                run.text = ""
        else:
            target_shape.text = new_text

        prs.save(file_path)
        return {
            "success": True,
            "file": file_path,
            "slide": slide_index,
            "placeholder": placeholder_name,
            "previous_value": previous,
            "new_value": new_text,
            "backup": backup
        }