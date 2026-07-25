"""Functional tests for individual tool wrappers — each just checks the
wrapper reshapes a step correctly and dispatches to the real underlying
component/handler (already unit-tested on its own)."""
import os

import openpyxl
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from tools.correct_document_tool import CorrectDocumentTool
from tools.docx_tool import DocxTool
from tools.organize_files_tool import OrganizeFilesTool
from tools.pdf_tool import PDFTool
from tools.pptx_tool import PPTXTool
from tools.sub_agent.browser_agent_tool import BrowserAgentTool
from tools.xlsx_tool import XLSXTool


def test_xlsx_tool_write_then_read(tmp_path):
    xlsx_path = str(tmp_path / "kpi_tracker.xlsx")
    wb = openpyxl.Workbook()
    wb.active["B7"] = "1.8M"
    wb.save(xlsx_path)

    tool = XLSXTool()
    write_result = tool.run(
        {"step_id": 1, "type": "write", "target": {"file": xlsx_path, "cell": "B7"}, "value": "2.4M"}
    )
    assert write_result["success"] is True
    assert write_result["previous_value"] == "1.8M"

    read_result = tool.run({"step_id": 2, "type": "read", "target": {"file": xlsx_path, "cell": "B7"}})
    assert read_result["value"] == "2.4M"


def test_pdf_tool_is_read_only(tmp_path):
    import fitz

    pdf_path = str(tmp_path / "report.pdf")
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Q3 revenue total: 2.4M")
    doc.save(pdf_path)
    doc.close()

    tool = PDFTool()
    result = tool.run({"step_id": 1, "type": "read", "target": {"file": pdf_path, "query": "revenue total"}})
    assert "2.4M" in result["value"]

    with pytest.raises(ValueError):
        tool.run({"step_id": 2, "type": "write", "target": {"file": pdf_path}, "value": "x"})


def test_pptx_tool_write_then_read(tmp_path):
    pptx_path = str(tmp_path / "deck.pptx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.name = "churn_value"
    box.text_frame.text = "2.9%"
    prs.save(pptx_path)

    tool = PPTXTool()
    write_result = tool.run(
        {
            "step_id": 1,
            "type": "write",
            "target": {"file": pptx_path, "slide": 1, "placeholder": "churn_value"},
            "value": "3.1%",
        }
    )
    assert write_result["success"] is True

    read_result = tool.run(
        {"step_id": 2, "type": "read", "target": {"file": pptx_path, "slide": 1, "placeholder": "churn_value"}}
    )
    assert read_result["value"] == "3.1%"


def test_docx_tool_write_then_read_paragraph(tmp_path):
    docx_path = str(tmp_path / "report.docx")
    doc = Document()
    doc.add_paragraph("Old text")
    doc.save(docx_path)

    tool = DocxTool()
    write_result = tool.run(
        {"step_id": 1, "type": "write", "target": {"file": docx_path, "paragraph_index": 0}, "value": "New text"}
    )
    assert write_result["success"] is True

    read_result = tool.run({"step_id": 2, "type": "read", "target": {"file": docx_path, "paragraph_index": 0}})
    assert read_result["value"] == "New text"


def test_organize_files_tool(tmp_path):
    pdf_path = tmp_path / "report.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\n%%EOF")
    dest_root = str(tmp_path / "organised")

    tool = OrganizeFilesTool()
    result = tool.run(
        {
            "step_id": 1,
            "type": "write",
            "target": {"files": [str(pdf_path)], "prompt": "organise by type", "dest_root": dest_root},
        }
    )
    assert result["success"] is True


def test_correct_document_tool(tmp_path):
    docx_path = str(tmp_path / "report.docx")
    doc = Document()
    doc.add_paragraph("This is teh report.")
    doc.save(docx_path)

    tool = CorrectDocumentTool()
    result = tool.run({"step_id": 1, "type": "write", "target": {"file": docx_path}})
    assert result["success"] is True
    assert os.path.exists(result["output_path"])


def test_browser_agent_tool_requires_instruction():
    tool = BrowserAgentTool()
    with pytest.raises(ValueError):
        tool.run({"step_id": 1, "type": "write", "target": {}})


def test_browser_agent_tool_wraps_run_result(monkeypatch):
    import tools.sub_agent.browser_agent_tool as mod

    monkeypatch.setattr(mod, "run_browser_agent", lambda instruction, **kw: {"result": "done: " + instruction, "step_count": 3})

    tool = BrowserAgentTool()
    result = tool.run({"step_id": 1, "type": "write", "target": {"instruction": "search for loop engineering"}})
    assert result == {"success": True, "value": "done: search for loop engineering", "step_count": 3}


def test_pptx_tool_create_with_explicit_slides_skips_llm(tmp_path, monkeypatch):
    import tools.pptx_tool as mod

    monkeypatch.setattr(mod, "structured_completion", lambda *_: (_ for _ in ()).throw(AssertionError("should not be called")))
    monkeypatch.setattr(mod.workspace, "WORKSPACE_DIR", str(tmp_path))

    tool = PPTXTool()
    result = tool.run(
        {
            "step_id": 1,
            "type": "create",
            "target": {
                "file": "explicit.pptx",
                "title": "T",
                "slides": [{"heading": "H", "bullets": ["b1"]}],
            },
        }
    )
    assert result["success"] is True
    assert os.path.exists(os.path.join(str(tmp_path), "explicit.pptx"))


def test_pptx_tool_create_from_source_material_calls_llm_outline(tmp_path, monkeypatch):
    import tools.pptx_tool as mod

    mock_outline = {"title": "Modi News", "slides": [{"heading": "Headline", "bullets": ["point one"]}]}
    monkeypatch.setattr(mod, "structured_completion", lambda *_: mock_outline)
    monkeypatch.setattr(mod.workspace, "WORKSPACE_DIR", str(tmp_path))

    tool = PPTXTool()
    result = tool.run(
        {
            "step_id": 2,
            "type": "create",
            "target": {"file": "from_search.pptx", "topic": "Modi"},
            "value": "raw search result text about Modi",
        }
    )
    assert result["success"] is True

    out_path = os.path.join(str(tmp_path), "from_search.pptx")
    prs = Presentation(out_path)
    assert prs.slides[0].shapes.title.text == "Modi News"


def test_pptx_tool_create_without_slides_or_source_raises():
    tool = PPTXTool()
    with pytest.raises(ValueError):
        tool.run({"step_id": 1, "type": "create", "target": {"file": "x.pptx"}})


def test_xlsx_tool_create_from_source_material_calls_llm_rows(tmp_path, monkeypatch):
    import tools.xlsx_tool as mod

    mock_outline = {"sheet_name": "Data", "rows": [["Metric", "Value"], ["Revenue", "2.4M"]]}
    monkeypatch.setattr(mod, "structured_completion", lambda *_: mock_outline)
    monkeypatch.setattr(mod.workspace, "WORKSPACE_DIR", str(tmp_path))

    tool = XLSXTool()
    result = tool.run(
        {
            "step_id": 1,
            "type": "create",
            "target": {"file": "from_search.xlsx", "topic": "revenue"},
            "value": "raw text with revenue figures",
        }
    )
    assert result["success"] is True

    out_path = os.path.join(str(tmp_path), "from_search.xlsx")
    wb = openpyxl.load_workbook(out_path)
    assert wb.active["A1"].value == "Metric"


def test_docx_tool_create_splits_source_text_into_paragraphs_without_llm(tmp_path, monkeypatch):
    import tools.docx_tool as mod

    monkeypatch.setattr(mod.workspace, "WORKSPACE_DIR", str(tmp_path))

    tool = DocxTool()
    result = tool.run(
        {
            "step_id": 1,
            "type": "create",
            "target": {"file": "from_search.docx", "title": "Modi"},
            "value": "First paragraph.\nSecond paragraph.",
        }
    )
    assert result["success"] is True

    out_path = os.path.join(str(tmp_path), "from_search.docx")
    doc = Document(out_path)
    texts = [p.text for p in doc.paragraphs]
    assert "First paragraph." in texts
    assert "Second paragraph." in texts


def test_create_with_bare_filename_resolves_into_workspace_dir(tmp_path, monkeypatch):
    monkeypatch.setattr("tools.workspace.WORKSPACE_DIR", str(tmp_path / "ws"))

    tool = PPTXTool()
    result = tool.run(
        {
            "step_id": 1,
            "type": "create",
            "target": {"file": "rr.pptx", "title": "T", "slides": [{"heading": "H", "bullets": ["b"]}]},
        }
    )
    assert result["file"] == os.path.join(str(tmp_path / "ws"), "rr.pptx")
    assert os.path.exists(result["file"])
