"""Standalone tests for the DOCX Adapter (component 18).

Run in isolation: pytest components/c18_docx_adapter
"""
import pytest
from docx import Document

from components.c18_docx_adapter.src.docx_adapter import read_or_write


def _make_document(path: str) -> None:
    doc = Document()
    doc.add_paragraph("Q3 revenue grew 12%.")
    doc.add_paragraph("Second paragraph.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[1].cells[0].text = "Revenue"
    table.rows[1].cells[1].text = "1.8M"
    doc.save(path)


def test_read_whole_document(tmp_path):
    docx_path = str(tmp_path / "report.docx")
    _make_document(docx_path)

    result = read_or_write({"file": docx_path, "operation": "read"})

    assert result["paragraphs"][0]["text"] == "Q3 revenue grew 12%."
    assert {"table_index": 0, "row": 1, "col": 1, "text": "1.8M"} in result["table_cells"]


def test_write_then_read_paragraph_round_trip(tmp_path):
    docx_path = str(tmp_path / "report.docx")
    _make_document(docx_path)

    write_result = read_or_write(
        {"file": docx_path, "operation": "write", "paragraph_index": 0, "value": "Q3 revenue grew 24%."}
    )
    assert write_result["success"] is True
    assert write_result["previous_value"] == "Q3 revenue grew 12%."

    read_result = read_or_write({"file": docx_path, "operation": "read", "paragraph_index": 0})
    assert read_result["value"] == "Q3 revenue grew 24%."


def test_write_then_read_table_cell_round_trip(tmp_path):
    docx_path = str(tmp_path / "report.docx")
    _make_document(docx_path)

    write_result = read_or_write(
        {"file": docx_path, "operation": "write", "table_index": 0, "row": 1, "col": 1, "value": "2.4M"}
    )
    assert write_result["success"] is True
    assert write_result["previous_value"] == "1.8M"

    read_result = read_or_write({"file": docx_path, "operation": "read", "table_index": 0, "row": 1, "col": 1})
    assert read_result["value"] == "2.4M"


def test_out_of_range_paragraph_errors_cleanly(tmp_path):
    docx_path = str(tmp_path / "report.docx")
    _make_document(docx_path)

    with pytest.raises(IndexError):
        read_or_write({"file": docx_path, "operation": "write", "paragraph_index": 99, "value": "x"})


def test_create_builds_a_new_document_with_title_and_paragraphs(tmp_path):
    docx_path = str(tmp_path / "new_doc.docx")

    result = read_or_write(
        {
            "file": docx_path,
            "operation": "create",
            "title": "Modi: Recent Developments",
            "paragraphs": ["First point.", "Second point."],
        }
    )

    assert result == {"success": True, "file": docx_path, "paragraph_count": 2}

    doc = Document(docx_path)
    all_text = [p.text for p in doc.paragraphs]
    assert "Modi: Recent Developments" in all_text[0]
    assert "First point." in all_text
    assert "Second point." in all_text
